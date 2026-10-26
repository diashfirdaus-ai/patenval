import io
from typing import Dict, Any, List
import pymupdf as fitz  # PyMuPDF
import docx


def build_demo_patent_pdf_bytes() -> bytes:
    """Generates a 3-page machine-readable PDF technical disclosure & patent specification using PyMuPDF."""
    pdf = fitz.open()

    p1 = pdf.new_page()
    p1.insert_text(
        (50, 60),
        "SYNTHETIC DEMONSTRATION DOCUMENT — ITENAS RESEARCH PORTFOLIO\n"
        "Document Type: Patent Specification & Technical Disclosure\n"
        "Patent Number: IDP000089421 (Granted in Indonesia, PDKI)\n"
        "Technology Title: Modular Photocatalytic Membrane Reactor for Industrial Textile Wastewater Reclamation\n\n"
        "1. Technical Architecture & Novelty (Page 1)\n"
        "The disclosed invention integrates a TiO2-graphene nanocomposite coating immobilized on a tubular ceramic "
        "ultrafiltration membrane. Laboratory and pilot tests confirm 94.2% chemical oxygen demand (COD) removal "
        "and 97.5% reactive azo dye degradation under low-energy UV-A LED irradiation at 365 nm.\n"
        "Unlike conventional coagulation-flocculation processes that generate hazardous chemical sludge, "
        "this photocatalytic reactor operates in a continuous closed-loop configuration with zero secondary sludge.",
        fontsize=10,
    )

    p2 = pdf.new_page()
    p2.insert_text(
        (50, 60),
        "2. Technology Readiness Level (TRL) & Subsystem Validation (Page 2)\n"
        "The subsystem prototype has been demonstrated in a relevant industrial environment at TRL 6 exit criteria.\n"
        "Pilot campaign conducted over 450 cumulative operating hours treating 15 m3/day of raw dyeing effluent.\n"
        "Engineering scaling parameters, transmembrane pressure profiles (1.8 bar nominal), and automated backflush "
        "cycles were logged and verified against NASA TRL 6 exit criteria adapted for process engineering systems.\n"
        "Specific energy consumption measured at 1.45 kWh/m3 of reclaimed process water.",
        fontsize=10,
    )

    p3 = pdf.new_page()
    p3.insert_text(
        (50, 60),
        "3. Intellectual Property Claims & Manufacturing Readiness (Page 3)\n"
        "Independent Claim 1 covers the thermal-sol-gel deposition method bonding the catalytic layer onto alumina "
        "supports without pore occlusion. Family jurisdictions include Indonesia (ID) and PCT (WO/2025/041822).\n"
        "Domestic supply chain audit confirms 82% local component availability (TKDN readiness) with two ISO-9001 "
        "ceramic fabricators in West Java qualified for pilot batch production.",
        fontsize=10,
    )

    out = pdf.tobytes()
    pdf.close()
    return out


def build_demo_scanned_pdf_bytes() -> bytes:
    """Generates a 1-page image-only scanned PDF with no text layer to verify OCR detection flags."""
    pdf = fitz.open()
    page = pdf.new_page()
    # Draw a simple raster image without inserting any text
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 120, 80), False)
    pix.set_rect(pix.irect, (220, 220, 220))
    page.insert_image(fitz.Rect(50, 50, 250, 180), pixmap=pix)
    out = pdf.tobytes()
    pdf.close()
    return out


def build_demo_market_docx_bytes() -> bytes:
    """Generates a multi-paragraph and table DOCX market & commercialization study using python-docx."""
    d = docx.Document()
    d.add_heading("SYNTHETIC DEMONSTRATION — Market & Commercialization Study", level=1)
    d.add_paragraph(
        "Target Market Segment: Indonesian textile dyeing and finishing clusters in Bandung, Majalaya, "
        "and Pekalongan facing strict Citarum Harum and KLHK wastewater discharge regulations."
    )
    d.add_paragraph(
        "Commercialization Pathway: Non-exclusive industrial equipment licensing combined with membrane "
        "cartridge replacement service agreements. Initial pilot licensee letter of intent signed for 3 units."
    )
    d.add_paragraph(
        "Regulatory & Standards Readiness: Effluent quality meets Ministry of Environment Regulation "
        "PermenLHK No. P.16/2019 threshold parameters for textile industry wastewater reuse."
    )

    table = d.add_table(rows=3, cols=3)
    headers = ["Deployment Phase", "Annual Unit Capacity", "Projected EBITDA (IDR)"]
    for idx, h in enumerate(headers):
        table.rows[0].cells[idx].text = h
    row1 = ["Year 1 Pilot Rollout", "4 Reactor Skids", "Rp 650.000.000"]
    row2 = ["Year 3 Commercial Scale", "12 Reactor Skids", "Rp 1.650.000.000"]
    for idx, val in enumerate(row1):
        table.rows[1].cells[idx].text = val
    for idx, val in enumerate(row2):
        table.rows[2].cells[idx].text = val

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_synthetic_evaluation_dataset_items() -> List[Dict[str, Any]]:
    """
    Constructs a matched 4-condition synthetic evaluation dataset comparing:
    - `human_expert`
    - `llm_only`
    - `multi_agent`
    - `multi_agent_validated`
    Prominently isolated and labeled as a synthetic demonstration fixture.
    """
    return [
        {
            "item_id": "SYNTH-TECH-01",
            "title": "Photocatalytic Membrane Reactor (Synthetic Benchmark Item 1)",
            "human_expert": {
                "indicator_scores": [4, 4, 3, 4, 4, 4, 3, 4, 3, 4],
                "valuation_idr": 962_450_000.0,
                "extractions": [
                    {"indicator_code": "TECH_NOVELTY", "location_ref": "page:1", "resolves_to_source": True},
                    {"indicator_code": "TECH_VALIDATION", "location_ref": "page:2", "resolves_to_source": True},
                    {"indicator_code": "TECH_PATENT_STRENGTH", "location_ref": "page:3", "resolves_to_source": True},
                    {"indicator_code": "COMM_MARKET_DEMAND", "location_ref": "para:2", "resolves_to_source": True},
                ],
            },
            "llm_only": {
                "indicator_scores": [5, 4, 4, 3, 5, 4, 4, 3, 4, 5],
                "valuation_idr": 1_240_000_000.0,
                "extractions": [
                    {"indicator_code": "TECH_NOVELTY", "location_ref": "page:1", "resolves_to_source": True},
                    {"indicator_code": "TECH_VALIDATION", "location_ref": "page:99", "resolves_to_source": False, "rejected_by_validator": False},
                    {"indicator_code": "COMM_MARKET_DEMAND", "location_ref": "para:2", "resolves_to_source": True},
                ],
            },
            "multi_agent": {
                "indicator_scores": [4, 4, 3, 4, 4, 4, 4, 4, 3, 4],
                "valuation_idr": 995_000_000.0,
                "extractions": [
                    {"indicator_code": "TECH_NOVELTY", "location_ref": "page:1", "resolves_to_source": True},
                    {"indicator_code": "TECH_VALIDATION", "location_ref": "page:2", "resolves_to_source": True},
                    {"indicator_code": "TECH_PATENT_STRENGTH", "location_ref": "page:3", "resolves_to_source": True},
                    {"indicator_code": "COMM_MARKET_DEMAND", "location_ref": "para:88", "resolves_to_source": False, "rejected_by_validator": False},
                ],
            },
            "multi_agent_validated": {
                "indicator_scores": [4, 4, 3, 4, 4, 4, 3, 4, 3, 4],
                "valuation_idr": 962_450_000.0,
                "extractions": [
                    {"indicator_code": "TECH_NOVELTY", "location_ref": "page:1", "resolves_to_source": True},
                    {"indicator_code": "TECH_VALIDATION", "location_ref": "page:2", "resolves_to_source": True},
                    {"indicator_code": "TECH_PATENT_STRENGTH", "location_ref": "page:3", "resolves_to_source": True},
                    {"indicator_code": "COMM_MARKET_DEMAND", "location_ref": "para:2", "resolves_to_source": True},
                    {"indicator_code": "HALLUCINATED_CLAIM", "location_ref": "page:99", "resolves_to_source": False, "rejected_by_validator": True},
                ],
            },
        },
        {
            "item_id": "SYNTH-TECH-02",
            "title": "Biopolymer Slow-Release Fertilizer Coating (Synthetic Benchmark Item 2)",
            "human_expert": {
                "indicator_scores": [3, 3, 3, 4, 3, 3, 3, 2, 3, 3],
                "valuation_idr": 610_000_000.0,
                "extractions": [
                    {"indicator_code": "TECH_NOVELTY", "location_ref": "page:1", "resolves_to_source": True},
                    {"indicator_code": "TECH_VALIDATION", "location_ref": "page:2", "resolves_to_source": True},
                    {"indicator_code": "COMM_REGULATORY", "location_ref": "para:4", "resolves_to_source": True},
                ],
            },
            "llm_only": {
                "indicator_scores": [4, 4, 3, 5, 4, 4, 2, 3, 4, 4],
                "valuation_idr": 820_000_000.0,
                "extractions": [
                    {"indicator_code": "TECH_NOVELTY", "location_ref": "page:1", "resolves_to_source": True},
                    {"indicator_code": "TECH_VALIDATION", "location_ref": "page:12", "resolves_to_source": False, "rejected_by_validator": False},
                ],
            },
            "multi_agent": {
                "indicator_scores": [3, 4, 3, 4, 3, 3, 3, 3, 3, 3],
                "valuation_idr": 640_000_000.0,
                "extractions": [
                    {"indicator_code": "TECH_NOVELTY", "location_ref": "page:1", "resolves_to_source": True},
                    {"indicator_code": "TECH_VALIDATION", "location_ref": "page:2", "resolves_to_source": True},
                    {"indicator_code": "COMM_REGULATORY", "location_ref": "para:4", "resolves_to_source": True},
                ],
            },
            "multi_agent_validated": {
                "indicator_scores": [3, 3, 3, 4, 3, 3, 3, 2, 3, 3],
                "valuation_idr": 615_000_000.0,
                "extractions": [
                    {"indicator_code": "TECH_NOVELTY", "location_ref": "page:1", "resolves_to_source": True},
                    {"indicator_code": "TECH_VALIDATION", "location_ref": "page:2", "resolves_to_source": True},
                    {"indicator_code": "COMM_REGULATORY", "location_ref": "para:4", "resolves_to_source": True},
                ],
            },
        },
        {
            "item_id": "SYNTH-TECH-03",
            "title": "IoT Edge Vibration Predictive Maintenance Sensor (Synthetic Benchmark Item 3)",
            "human_expert": {
                "indicator_scores": [4, 5, 4, 3, 4, 4, 4, 4, 4, 3],
                "valuation_idr": 1_350_000_000.0,
                "extractions": [
                    {"indicator_code": "TECH_NOVELTY", "location_ref": "page:1", "resolves_to_source": True},
                    {"indicator_code": "TECH_SCALABILITY", "location_ref": "table:1", "resolves_to_source": True},
                ],
            },
            "llm_only": {
                "indicator_scores": [3, 4, 3, 2, 3, 3, 3, 3, 3, 2],
                "valuation_idr": 980_000_000.0,
                "extractions": [
                    {"indicator_code": "TECH_NOVELTY", "location_ref": "page:1", "resolves_to_source": True},
                ],
            },
            "multi_agent": {
                "indicator_scores": [4, 5, 4, 3, 4, 4, 4, 3, 4, 3],
                "valuation_idr": 1_380_000_000.0,
                "extractions": [
                    {"indicator_code": "TECH_NOVELTY", "location_ref": "page:1", "resolves_to_source": True},
                    {"indicator_code": "TECH_SCALABILITY", "location_ref": "table:1", "resolves_to_source": True},
                ],
            },
            "multi_agent_validated": {
                "indicator_scores": [4, 5, 4, 3, 4, 4, 4, 4, 4, 3],
                "valuation_idr": 1_355_000_000.0,
                "extractions": [
                    {"indicator_code": "TECH_NOVELTY", "location_ref": "page:1", "resolves_to_source": True},
                    {"indicator_code": "TECH_SCALABILITY", "location_ref": "table:1", "resolves_to_source": True},
                ],
            },
        },
    ]
